"""Escritura de los estados de cuenta a libros de Excel (openpyxl)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .parser import EstadoCuenta

INK = "10233A"
ICE = "DDEEFB"
FMT_NUM = "#,##0.00;[Red]-#,##0.00"
FMT_INT = "0"
FMT_DATE = "DD/MM/YYYY"
FMT_DATETIME = "DD/MM/YYYY HH:MM:SS"

_FUENTE_CAB = Font(bold=True, color="FFFFFF")
_RELLENO_CAB = PatternFill("solid", fgColor=INK)
_FUENTE_TOTAL = Font(bold=True)
_RELLENO_TOTAL = PatternFill("solid", fgColor=ICE)
_FUENTE_TITULO = Font(bold=True, size=16)
_FUENTE_ETIQUETA = Font(bold=True)

PREFIJO_CAB = ("Archivo", "Cuenta", "Desde", "Hasta")


def _prefijo(ec: EstadoCuenta, combinado: bool) -> list:
    return [ec.archivo.name, ec.cuenta, ec.periodo_inicio, ec.periodo_fin] if combinado else []


def _ancho(valor) -> int:
    if valor is None:
        return 0
    if isinstance(valor, datetime):
        return 19
    if isinstance(valor, date):
        return 10
    if isinstance(valor, float):
        return len(f"{valor:,.2f}") + 1
    return len(str(valor))


def _escribir(ws: Worksheet, fila: int, col: int, valor, fuente: Font | None = None, relleno: PatternFill | None = None):
    celda = ws.cell(row=fila, column=col, value=valor)
    if isinstance(valor, datetime):
        celda.number_format = FMT_DATETIME
    elif isinstance(valor, date):
        celda.number_format = FMT_DATE
        celda.alignment = Alignment(horizontal="center")
    elif isinstance(valor, float):
        celda.number_format = FMT_NUM
    elif isinstance(valor, int):
        celda.number_format = FMT_INT
        celda.alignment = Alignment(horizontal="center")
    if fuente:
        celda.font = fuente
    if relleno:
        celda.fill = relleno
    return celda


class _Tabla:
    """Escribe una tabla con encabezado, filas de datos y filas de totales, y ajusta anchos."""

    def __init__(self, ws: Worksheet, encabezados: list[str], fila: int = 1):
        self.ws = ws
        self.fila_cab = fila
        self.fila = fila + 1
        self.n = len(encabezados)
        self.anchos = [len(h) for h in encabezados]
        for i, h in enumerate(encabezados, start=1):
            c = _escribir(ws, fila, i, h, _FUENTE_CAB, _RELLENO_CAB)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def fila_datos(self, valores: list, total: bool = False) -> None:
        for i, v in enumerate(valores, start=1):
            _escribir(self.ws, self.fila, i, v, _FUENTE_TOTAL if total else None, _RELLENO_TOTAL if total else None)
            self.anchos[i - 1] = max(self.anchos[i - 1], _ancho(v))
        if total:  # rellena también las celdas vacías de la fila de totales
            for i in range(len(valores) + 1, self.n + 1):
                self.ws.cell(row=self.fila, column=i).fill = _RELLENO_TOTAL
        self.fila += 1

    def cerrar(self) -> None:
        ws = self.ws
        for i, ancho in enumerate(self.anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 9), 48)
        ws.freeze_panes = ws.cell(row=self.fila_cab + 1, column=1)
        if self.fila - 1 > self.fila_cab:
            ws.auto_filter.ref = f"A{self.fila_cab}:{get_column_letter(self.n)}{self.fila - 1}"


# ----------------------------------------------------------------------------- hojas

def _hoja_resumen(wb: Workbook, estados: list[EstadoCuenta], combinado: bool) -> None:
    ws = wb.active
    ws.title = "Resumen"
    if combinado:
        t = _Tabla(ws, [*PREFIJO_CAB, "Razón social", "Sociedad", "Concepto", "Fecha / Cantidad", "Importe"])
        for ec in estados:
            for r in ec.resumen:
                t.fila_datos([*_prefijo(ec, True), ec.razon_social, ec.sociedad, r.concepto, r.dato, r.importe])
        t.cerrar()
        return

    ec = estados[0]
    _escribir(ws, 1, 1, "Estado de cuenta", _FUENTE_TITULO)
    datos = (
        ("Cuenta", ec.cuenta), ("Razón social", ec.razon_social), ("Sociedad", ec.sociedad),
        ("Periodo desde", ec.periodo_inicio), ("Periodo hasta", ec.periodo_fin),
        ("Fecha de emisión", ec.emision), ("Hojas", ec.hojas), ("Archivo", ec.archivo.name),
    )
    fila = 3
    for etiqueta, valor in datos:
        _escribir(ws, fila, 1, etiqueta, _FUENTE_ETIQUETA)
        _escribir(ws, fila, 2, valor)
        fila += 1
    t = _Tabla(ws, ["Concepto", "Fecha / Cantidad", "Importe"], fila + 1)
    for r in ec.resumen:
        t.fila_datos([r.concepto, r.dato, r.importe])
    t.cerrar()
    ws.freeze_panes = None
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 30)


def _hoja_movimientos(wb: Workbook, estados: list[EstadoCuenta], combinado: bool) -> None:
    ws = wb.create_sheet("Movimientos")
    cab = ["Fecha", "Tipo", "Referencia", "Asignación", "Cargos", "Abonos", "Saldo final",
           "Pago", "Notas de crédito", "Documento de aplicación"]
    t = _Tabla(ws, [*(PREFIJO_CAB if combinado else ()), *cab])
    for ec in estados:
        p = _prefijo(ec, combinado)
        for m in ec.movimientos:
            t.fila_datos([*p, m.fecha, m.tipo, m.referencia, m.asignacion, m.cargos, m.abonos, m.saldo_final, None, None, None])
            for a in m.aplicaciones:
                t.fila_datos([*p, None, "Aplicación", None, None, None, None, None, a.pago, a.nota_credito, a.documento])
        tot = ec.totales_movimientos
        if tot:
            t.fila_datos([*p, None, "CIFRAS CONTROL", None, None, tot.get("cargos"), tot.get("abonos"), None,
                          tot.get("pago"), tot.get("nota_credito"), None], total=True)
    t.cerrar()


def _hoja_aplicaciones(wb: Workbook, estados: list[EstadoCuenta], combinado: bool) -> None:
    ws = wb.create_sheet("Aplicaciones")
    cab = ["Fecha del pago", "Referencia del pago", "Asignación del pago", "Importe aplicado",
           "Nota de crédito", "Documento aplicado"]
    t = _Tabla(ws, [*(PREFIJO_CAB if combinado else ()), *cab])
    for ec in estados:
        p = _prefijo(ec, combinado)
        for m in ec.movimientos:
            for a in m.aplicaciones:
                t.fila_datos([*p, m.fecha, m.referencia, m.asignacion, a.pago, a.nota_credito, a.documento])
    t.cerrar()


def _hoja_saldo(wb: Workbook, estados: list[EstadoCuenta], combinado: bool) -> None:
    ws = wb.create_sheet("Saldo final")
    cab = ["Fecha doc", "Vencimiento", "Referencia", "Asignación", "Cargo", "Abono", "Saldo", "Descripción"]
    t = _Tabla(ws, [*(PREFIJO_CAB if combinado else ()), *cab])
    for ec in estados:
        p = _prefijo(ec, combinado)
        for d in ec.saldo_final:
            t.fila_datos([*p, d.fecha_doc, d.vencimiento, d.referencia, d.asignacion, d.cargo, d.abono, d.saldo, d.descripcion])
        tot = ec.totales_saldo
        if tot:
            t.fila_datos([*p, None, None, "CIFRAS CONTROL", None, tot.get("cargo"), tot.get("abono"), None, None], total=True)
    t.cerrar()


def _hoja_avisos(wb: Workbook, estados: list[EstadoCuenta]) -> None:
    avisos = [(ec.archivo.name, a) for ec in estados for a in ec.avisos]
    if not avisos:
        return
    ws = wb.create_sheet("Avisos")
    t = _Tabla(ws, ["Archivo", "Aviso"])
    for fila in avisos:
        t.fila_datos(list(fila))
    t.cerrar()


# ----------------------------------------------------------------------------- API

def guardar(estados: list[EstadoCuenta], ruta: str | Path, combinado: bool = False) -> Path:
    """Escribe uno o varios estados de cuenta en un libro. Con `combinado`, cada hoja lleva
    las columnas Archivo/Cuenta/Desde/Hasta para distinguir de qué estado viene cada fila."""
    if not estados:
        raise ValueError("No hay estados de cuenta que guardar")
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    estados = sorted(estados, key=lambda e: (e.periodo_inicio or date.min, e.archivo.name))
    wb = Workbook()
    _hoja_resumen(wb, estados, combinado)
    _hoja_movimientos(wb, estados, combinado)
    _hoja_aplicaciones(wb, estados, combinado)
    _hoja_saldo(wb, estados, combinado)
    _hoja_avisos(wb, estados)
    wb.save(str(ruta))
    return ruta


def guardar_generico(hojas: list[tuple[str, list[list[str]]]], ruta: str | Path) -> Path:
    """Vuelca las tablas/texto de un PDF de formato desconocido, una hoja por tabla."""
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    usados: set[str] = set()
    for nombre, filas in hojas or [("Vacío", [])]:
        nombre = nombre[:31]
        base, k = nombre, 2
        while nombre in usados:
            nombre, k = f"{base[:28]} {k}", k + 1
        usados.add(nombre)
        ws = wb.create_sheet(nombre)
        anchos: dict[int, int] = {}
        for i, fila in enumerate(filas, start=1):
            for j, valor in enumerate(fila, start=1):
                ws.cell(row=i, column=j, value=valor)
                anchos[j] = max(anchos.get(j, 0), len(valor or ""))
        for j, ancho in anchos.items():
            ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 9), 60)
    wb.save(str(ruta))
    return ruta
